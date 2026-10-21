











from __future__ import annotations

import csv
import json
import logging
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)


MAX_TEXT_CHARS = 120_000

MAX_TABLE_ROWS = 200
MAX_TABLE_SHEETS = 5

MAX_PDF_PAGES = 20

MAX_CELL_CHARS = 500

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".ico"}
MARKDOWN_EXTS = {".md", ".markdown"}
PLAIN_TEXT_EXTS = {".txt", ".log", ".yaml", ".yml", ".ini", ".conf", ".py", ".ts", ".js"}
SPREADSHEET_EXTS = {".xlsx", ".xlsm", ".xltx", ".xltm"}


def _cell(value: Any) -> Any:

    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    return text[:MAX_CELL_CHARS] + ("…" if len(text) > MAX_CELL_CHARS else "")


def _read_text(path: Path) -> tuple[str, bool]:

    try:
        with open(path, encoding="utf-8", errors="replace") as fh:
            text = fh.read(MAX_TEXT_CHARS + 1)
    except OSError as exc:
        raise PreviewError(f"无法读取文件：{exc}") from exc
    if len(text) > MAX_TEXT_CHARS:
        return text[:MAX_TEXT_CHARS], True
    return text, False





def _preview_svg(path: Path) -> dict[str, Any]:
    content, truncated = _read_text(path)
    return {"kind": "svg", "content": content, "truncated": truncated}


def _preview_markdown(path: Path) -> dict[str, Any]:
    content, truncated = _read_text(path)
    return {"kind": "markdown", "content": content, "truncated": truncated}


def _preview_text(path: Path) -> dict[str, Any]:
    content, truncated = _read_text(path)
    return {"kind": "text", "content": content, "truncated": truncated}


def _preview_json(path: Path) -> dict[str, Any]:
    content, truncated = _read_text(path)
    try:
        parsed = json.loads(content)
    except json.JSONDecodeError as exc:
        
        return {"kind": "text", "content": content, "truncated": truncated, "note": f"JSON 解析失败：{exc}"}
    return {"kind": "json", "content": json.dumps(parsed, ensure_ascii=False, indent=2)[:MAX_TEXT_CHARS]}


def _preview_csv(path: Path) -> dict[str, Any]:

    rows: list[list[Any]] = []
    total = 0
    try:
        with open(path, encoding="utf-8", errors="replace", newline="") as fh:
            reader = csv.reader(fh)
            for index, row in enumerate(reader):
                total = index + 1
                if index > MAX_TABLE_ROWS:
                    continue
                rows.append([_cell(item) for item in row])
    except (OSError, csv.Error) as exc:
        raise PreviewError(f"CSV 解析失败：{exc}") from exc

    if not rows:
        return {"kind": "table", "sheets": [], "note": "文件是空的"}

    header, body = rows[0], rows[1:]
    return {
        "kind": "table",
        "sheets": [
            {
                "name": path.stem,
                "columns": [str(item) if item is not None else "" for item in header],
                "rows": body,
                "total_rows": max(total - 1, 0),
                "truncated": total - 1 > len(body),
            }
        ],
    }


def _preview_xlsx(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  
        raise PreviewError("未安装 openpyxl，无法预览 Excel。") from exc

    sheets: list[dict[str, Any]] = []
    try:
        
        
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  
        raise PreviewError(f"Excel 解析失败：{exc}") from exc

    try:
        for name in workbook.sheetnames[:MAX_TABLE_SHEETS]:
            sheet = workbook[name]
            rows_iter = sheet.iter_rows(values_only=True)
            header: list[str] = []
            body: list[list[Any]] = []
            total = 0
            for index, row in enumerate(rows_iter):
                if index == 0:
                    header = [str(item) if item is not None else "" for item in row]
                    continue
                total = index
                if len(body) < MAX_TABLE_ROWS:
                    body.append([_cell(item) for item in row])
            if not header:
                continue
            sheets.append(
                {
                    "name": name,
                    "columns": header,
                    "rows": body,
                    "total_rows": total,
                    "truncated": total > len(body),
                }
            )
    except Exception as exc:  
        raise PreviewError(f"Excel 读取失败：{exc}") from exc
    finally:
        workbook.close()

    if not sheets:
        return {"kind": "table", "sheets": [], "note": "工作簿里没有可预览的数据"}
    return {"kind": "table", "sheets": sheets}


def _preview_docx(path: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError as exc:  
        raise PreviewError("未安装 python-docx，无法预览 Word 文档。") from exc

    try:
        document = Document(str(path))
    except Exception as exc:  
        raise PreviewError(f"Word 解析失败：{exc}") from exc

    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append(f"{'#' * min(int(level), 6)} {text}")
        else:
            lines.append(text)

    for index, table in enumerate(document.tables, start=1):
        lines.append(f"\n**表格 {index}**\n")
        for row_index, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if row_index == 0:
                lines.append("|" + "---|" * len(cells))
            if row_index >= MAX_TABLE_ROWS:
                lines.append("| … |")
                break

    content = "\n".join(lines)
    truncated = len(content) > MAX_TEXT_CHARS
    return {"kind": "markdown", "content": content[:MAX_TEXT_CHARS], "truncated": truncated}


def _preview_pdf(path: Path) -> dict[str, Any]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  
        raise PreviewError("未安装 pypdf，无法预览 PDF。") from exc

    try:
        reader = PdfReader(str(path))
        pages: list[str] = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            pages.append(page.extract_text() or "")
    except Exception as exc:  
        raise PreviewError(f"PDF 解析失败：{exc}") from exc

    content = "\n\n".join(item.strip() for item in pages if item.strip())
    if not content:
        
        return {
            "kind": "unsupported",
            "reason": "这个 PDF 没有可提取的文字（可能是扫描件）。可以用「另存为」在本地打开。",
        }
    return {"kind": "text", "content": content[:MAX_TEXT_CHARS], "truncated": len(reader.pages) > MAX_PDF_PAGES}


class PreviewError(Exception):
    pass





def build_preview(
    path: Path,
    *,
    original_name: str,
    mime: str | None = None,
    download_url: str | None = None,
) -> dict[str, Any]:





    suffix = Path(original_name).suffix.lower()
    base: dict[str, Any] = {
        "name": original_name,
        "suffix": suffix,
        "size": path.stat().st_size if path.is_file() else 0,
        "mime": mime,
        "download_url": download_url,
    }

    if not path.is_file():
        return {**base, "kind": "unsupported", "reason": "文件已丢失。"}

    try:
        if suffix in IMAGE_EXTS:
            payload: dict[str, Any] = {
                "kind": "image",
                
                "url": download_url,
            }
        elif suffix == ".svg":
            payload = _preview_svg(path)
        elif suffix in MARKDOWN_EXTS:
            payload = _preview_markdown(path)
        elif suffix == ".csv":
            payload = _preview_csv(path)
        elif suffix in SPREADSHEET_EXTS:
            payload = _preview_xlsx(path)
        elif suffix == ".docx":
            payload = _preview_docx(path)
        elif suffix == ".pdf":
            payload = _preview_pdf(path)
        elif suffix == ".json":
            payload = _preview_json(path)
        elif suffix in PLAIN_TEXT_EXTS:
            payload = _preview_text(path)
        else:
            payload = {
                "kind": "unsupported",
                "reason": f"暂不支持预览 {suffix or '该类型'} 文件，可以用「另存为」在本地打开。",
            }
    except PreviewError as exc:
        payload = {"kind": "unsupported", "reason": str(exc)}
    except Exception as exc:  
        logger.exception("预览文件失败: %s", original_name)
        payload = {"kind": "unsupported", "reason": f"预览失败：{type(exc).__name__}"}

    return {**base, **payload}


__all__ = ["MAX_TABLE_ROWS", "PreviewError", "build_preview"]
