#!/usr/bin/env python3
# -*- coding: utf-8 -*-




























import os

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from consult_docx_utils import set_run_font, add_para, setup_page, add_page_field




def write_cell(cell, text, bold=False, size=11, align_center=False, font="微软雅黑"):

    cell.text = ""
    p = cell.paragraphs[0]
    if align_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, line in enumerate(text.split("\n")):
        if i > 0:
            p = cell.add_paragraph()
            if align_center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        set_run_font(run, font, size, bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER




def build_minutes(data, output_path):

    doc = Document()

    
    section = setup_page(doc, top=3.0, bottom=2.5, left=3.0, right=3.0)

    
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp, font_name="微软雅黑", size=11, label_with_total=False)

    
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(data["标题"])
    set_run_font(run, "微软雅黑", 18, True)

    
    table = doc.add_table(rows=0, cols=7)
    table.style = 'Table Grid'

    def add_row():
        return table.add_row()

    
    r = add_row()
    write_cell(r.cells[0], "会议时间", bold=True, align_center=True)
    m = r.cells[1].merge(r.cells[4])
    write_cell(m, data["会议时间"])
    write_cell(r.cells[5], "会议地点", bold=True, align_center=True)
    write_cell(r.cells[6], data["会议地点"])

    
    r = add_row()
    write_cell(r.cells[0], "会议主持人", bold=True, align_center=True)
    m = r.cells[1].merge(r.cells[6])
    write_cell(m, data["主持人"])

    
    r = add_row()
    write_cell(r.cells[0], "与 会 者", bold=True, align_center=True)
    m = r.cells[1].merge(r.cells[6])
    write_cell(m, data["与会者"], bold=True)

    
    r = add_row()
    write_cell(r.cells[0], "会议议程", bold=True, align_center=True)
    m = r.cells[1].merge(r.cells[6])
    agenda_text = "\n".join(f"{i+1}、{a}" for i, a in enumerate(data["议程"]))
    write_cell(m, agenda_text)

    
    r = add_row()
    m = r.cells[0]
    for c in r.cells[1:]:
        m = m.merge(c)
    res_lines = ["会议决议："]
    for i, x in enumerate(data["决议"]):
        res_lines.append(f"{i+1}、{x}")
    if data.get("补充决议"):
        res_lines.append("补充决议：")
        for x in data["补充决议"]:
            res_lines.append(x)
    write_cell(m, "\n".join(res_lines), bold=True)

    
    r = add_row()
    m = r.cells[0].merge(r.cells[1])
    write_cell(m, "会议安排事项", bold=True, align_center=True)
    write_cell(r.cells[2], "责任人", bold=True, align_center=True)
    write_cell(r.cells[3], "追踪人", bold=True, align_center=True)
    write_cell(r.cells[4], "完成时限", bold=True, align_center=True)
    m2 = r.cells[5].merge(r.cells[6])
    write_cell(m2, "备注", bold=True, align_center=True)

    
    for act in data["行动项"]:
        r = add_row()
        m = r.cells[0].merge(r.cells[1])
        write_cell(m, act["内容"])
        write_cell(r.cells[2], act["责任人"])
        write_cell(r.cells[3], act["追踪人"])
        write_cell(r.cells[4], act["完成时限"])
        m2 = r.cells[5].merge(r.cells[6])
        write_cell(m2, act.get("备注", ""))

    
    r = add_row()
    for c in r.cells:
        write_cell(c, "")

    
    r = add_row()
    m = r.cells[0].merge(r.cells[1])
    write_cell(m, "整理人", bold=True, align_center=True)
    write_cell(r.cells[2], data["整理人"])
    write_cell(r.cells[3], "抄送", bold=True, align_center=True)
    m2 = r.cells[4].merge(r.cells[6])
    write_cell(m2, data["抄送"])

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path




if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(f"用法: python3 {sys.argv[0]} input.yaml output.docx")
        sys.exit(1)
    import yaml
    with open(sys.argv[1], encoding="utf-8") as f:
        data = yaml.safe_load(f)
    path = build_minutes(data, sys.argv[2])
    print(f"已生成: {path}")
