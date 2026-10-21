#!/usr/bin/env python3
# -*- coding: utf-8 -*-










from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_run_font(run, font_name, size=None, bold=False):

    run.font.name = font_name
    r = run._element
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        r.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    if size:
        run.font.size = Pt(size)
    run.bold = bool(bold)


def add_para(doc, text, font="仿宋", size=14, bold=False,
             align="left", space_after=0, indent_first=False):

    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    if indent_first:
        pf.first_line_indent = Pt(28)
    if text:
        run = p.add_run(text)
        set_run_font(run, font, size, bold)
    return p


def setup_page(doc, page_w=21.0, page_h=29.7,
               top=2.5, bottom=3.0, left=3.0, right=3.0):

    section = doc.sections[0]
    section.page_width = Cm(page_w)
    section.page_height = Cm(page_h)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    return section


def add_page_field(paragraph, font_name="仿宋", size=14, label_with_total=True):

    
    s1 = paragraph.add_run(" ")
    set_run_font(s1, font_name, size, False)

    
    r1 = paragraph.add_run()
    set_run_font(r1, font_name, size, False)
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    t1 = OxmlElement('w:instrText'); t1.text = 'PAGE'
    e1 = OxmlElement('w:fldChar'); e1.set(qn('w:fldCharType'), 'end')
    r1._r.append(f1); r1._r.append(t1); r1._r.append(e1)

    if label_with_total:
        sep = paragraph.add_run(" / ")
        set_run_font(sep, font_name, size, False)
        
        r2 = paragraph.add_run()
        set_run_font(r2, font_name, size, False)
        f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'begin')
        t2 = OxmlElement('w:instrText'); t2.text = 'NUMPAGES'
        e2 = OxmlElement('w:fldChar'); e2.set(qn('w:fldCharType'), 'end')
        r2._r.append(f2); r2._r.append(t2); r2._r.append(e2)

    
    s2 = paragraph.add_run(" ")
    set_run_font(s2, font_name, size, False)


def wrap_filename(name):

    if name.startswith("《"):
        return name
    if "：" in name and "《" not in name:
        return name
    return f"《{name}》"
