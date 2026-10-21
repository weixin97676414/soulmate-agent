#!/usr/bin/env python3





import sys
import re
import os
import subprocess
import tempfile
import shutil
from pathlib import Path

def _find_mermaid_script() -> str:

    import os as _os
    
    _search_dirs = []
    if '__file__' in dir():
        _search_dirs.append(str(Path(__file__).resolve().parent))
    
    _search_dirs.append(str(Path.cwd()))
    _search_dirs.append(str(Path.cwd().parent))  
    
    _search_dirs.extend([
        '/home/yhongm/.hermes/skills/software-manager-skill/scripts',
        '/home/yhongm/.claude/skills/software-manager-skill/scripts',
        str(Path.home() / '.hermes/skills/software-manager-skill/scripts'),
    ])
    
    for _user in ['yhong', 'yhongm']:
        _search_dirs.append(f'/mnt/c/Users/{_user}/.claude/skills/software-manager-skill/scripts')
        _search_dirs.append(f'/mnt/c/Users/{_user}/.hermes/skills/software-manager-skill/scripts')
    for _dir in _search_dirs:
        _candidate = _os.path.join(_dir, 'mermaid_render_multi.js')
        if _os.path.isfile(_candidate):
            return _candidate
    return ''

def render_mermaid_to_png(mermaid_code: str, output_path: str) -> bool:



    
    with tempfile.NamedTemporaryFile(mode='w', suffix='.mmd', delete=False, dir='/tmp') as f:
        f.write(mermaid_code)
        mmd_path = f.name

    try:
        
        mm_script = _find_mermaid_script()
        if not mm_script or not os.path.isfile(mm_script):
            print("ERROR: mermaid_render_multi.js not found in any search path.", file=sys.stderr)
            os.unlink(mmd_path)
            return False

        result = subprocess.run(
            ['node', mm_script, mmd_path, output_path],
            capture_output=True, timeout=60
        )
        return result.returncode == 0 and os.path.exists(output_path)
    finally:
        try:
            os.unlink(mmd_path)
        except:
            pass


def parse_markdown_sections(text: str) -> list:

    sections = []
    
    parts = re.split(r'(```mermaid\n[\s\S]*?```)', text)

    for part in parts:
        if part.startswith('```mermaid'):
            
            code = re.sub(r'^```mermaid\n?', '', part)
            code = re.sub(r'\n?```$', '', code)
            sections.append({'type': 'mermaid', 'content': code.strip()})
        elif part.strip():
            sections.append({'type': 'markdown', 'content': part})

    return sections


def heading_level(line: str) -> int:

    m = re.match(r'^(#{1,6})\s+', line)
    return len(m.group(1)) if m else 0


def export_md_to_docx(md_text: str, docx_path: str, title: str = "PRD Document"):

    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.style import WD_STYLE_TYPE
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        print("ERROR: python-docx not installed. Run: pip3 install python-docx", file=sys.stderr)
        return False

    
    img_dir = tempfile.mkdtemp(prefix='prd_mmd_')
    img_counter = [0]

    def next_img_path() -> str:
        img_counter[0] += 1
        return os.path.join(img_dir, f'mermaid_{img_counter[0]:03d}.png')

    doc = Document()

    
    title_para = doc.add_heading(title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    
    lines = md_text.split('\n')
    in_code_block = False
    code_block_content = []
    code_lang = ''

    def flush_code_block(doc, lang, content_lines):
        if not content_lines:
            return
        code_text = '\n'.join(content_lines)
        if lang == 'mermaid':
            
            img_path = next_img_path()
            ok = render_mermaid_to_png(code_text, img_path)
            if ok and os.path.exists(img_path):
                
                para = doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run()
                run.add_picture(img_path, width=Inches(5.5))
                doc.add_paragraph()  
            else:
                
                para = doc.add_paragraph()
                para.style = 'No Spacing'
                run = para.add_run('[Mermaid图表渲染失败，已降级为代码块]\n' + code_text)
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
        else:
            
            para = doc.add_paragraph()
            para.style = 'No Spacing'
            run = para.add_run(code_text)
            run.font.name = 'Courier New'
            run.font.size = Pt(9)

    for line in lines:
        if line.startswith('```'):
            if not in_code_block:
                
                in_code_block = True
                code_lang = line[3:].strip()
                code_block_content = []
            else:
                
                in_code_block = False
                flush_code_block(doc, code_lang, code_block_content)
                code_lang = ''
                code_block_content = []
        elif in_code_block:
            code_block_content.append(line)
        else:
            
            hlvl = heading_level(line)
            if hlvl > 0:
                
                heading_text = line.lstrip('#').strip()
                doc.add_heading(heading_text, level=hlvl)
            elif line.startswith('- ') or line.startswith('* '):
                
                doc.add_paragraph(line[2:], style='List Bullet')
            elif re.match(r'^\d+\.\s+', line):
                
                m = re.match(r'^(\d+\.)\s+(.*)', line)
                if m:
                    p = doc.add_paragraph(m.group(2), style='List Number')
                    p.style._element.set(qn('w:numId'), m.group(1))
            elif line.startswith('|'):
                
                
                doc.add_paragraph(line)
            elif line.strip() == '':
                doc.add_paragraph()
            else:
                
                inline_code = re.sub(r'`([^`]+)`', r'\1', line)
                doc.add_paragraph(inline_code)

    
    doc.save(docx_path)

    
    try:
        shutil.rmtree(img_dir)
    except:
        pass

    return True


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python3 prd_export.py <input.md> <output.docx> [title]")
        sys.exit(1)

    
    def normalize_path(p: str) -> str:

        import re
        
        cleaned = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', p)
        
        cleaned = re.sub(r'https?://\S+$', '', cleaned).strip()
        cleaned = re.sub(r'http://\S+$', '', cleaned).strip()
        
        if cleaned.startswith('/mnt/c/') or cleaned.startswith('C:\\') or cleaned.startswith('C:/'):
            
            pass
        elif len(cleaned) > 2 and cleaned[1] == ':':
            
            drive = cleaned[0].upper()
            rest = cleaned[2:].replace('\\', '/').replace(':', '')
            cleaned = f'/mnt/{drive.lower()}/{rest}'
        return cleaned

    md_path_raw = sys.argv[1]
    docx_path_raw = sys.argv[2]

    md_path = normalize_path(md_path_raw)
    docx_path = normalize_path(docx_path_raw)
    title = sys.argv[3] if len(sys.argv) > 3 else Path(md_path).stem

    if not os.path.exists(md_path):
        print(f"ERROR: Input file not found: {md_path}  (raw: {md_path_raw})", file=sys.stderr)
        sys.exit(1)

    md_text = Path(md_path).read_text(encoding='utf-8')
    success = export_md_to_docx(md_text, docx_path, title)

    if success:
        print(f"OK: {docx_path}")
    else:
        print(f"ERROR: Export failed", file=sys.stderr)
        sys.exit(1)
