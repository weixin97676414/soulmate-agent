# -*- coding: utf-8 -*-












import argparse
import datetime
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib_outline import (  
    load_json, save_json, material_index, die,
)

CN_FONT = "微软雅黑"


def _set_font(run, size=None, bold=None, color=None, italic=None, font=CN_FONT):
    from docx.shared import Pt, RGBColor
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rf is None:
        from docx.oxml import OxmlElement
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", font)
    rf.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii", font)
    rf.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi", font)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def _para(doc, text, size=11, bold=False, color=None, italic=False,
          space_after=6, align=None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph()
    r = p.add_run(text)
    _set_font(r, size=size, bold=bold, color=color, italic=italic)
    pf = p.paragraph_format
    pf.space_after = __import__("docx").shared.Pt(space_after)
    if align is not None:
        p.alignment = align
    return p


def reindex(ol):

    for i, p in enumerate(ol.get("pages", []), 1):
        p["id"] = "P-%d" % i
        for j, t in enumerate(p.get("points", []), 1):
            t["id"] = "T-%d-%d" % (i, j)
    return ol


def render_md(ol, mats):
    meta = ol.get("meta", {}) or {}
    L = ["# %s —— 演示大纲" % meta.get("name", "演示材料"), ""]
    L.append("- 受众：%s" % (meta.get("audience") or "（待填）"))
    L.append("- 演示目标：%s" % (meta.get("goal") or "（待填）"))
    L.append("- 计划时长：%s 分钟" % (meta.get("duration_min") or "（待填）"))
    L.append("- 生成时间：%s" % (meta.get("created_at") or ""))
    if meta.get("edited_by"):
        L.append("- 人工调整：%s（%s）" % (meta["edited_by"], meta.get("edit_note", "")))
    L.append("")
    L.append("## 页面结构")
    L.append("")
    L.append("| 页码 | 角色 | 标题 | 要点数 | 时长(分) | 素材来源 |")
    L.append("|---|---|---|---|---|---|")
    for p in ol.get("pages", []):
        sids = []
        for t in p.get("points", []):
            for s in (t.get("sids") or []):
                if s not in sids:
                    sids.append(s)
        L.append("| %s | %s | %s | %d | %s | %s |"
                 % (p.get("id"), p.get("role", ""), p.get("title", ""),
                    len(p.get("points", [])), p.get("minutes", ""),
                    "、".join(sids) or "-"))
    L.append("")
    for p in ol.get("pages", []):
        L.append("## %s %s" % (p.get("id"), p.get("title", "")))
        L.append("")
        if p.get("goal"):
            L.append("**目标**：%s" % p["goal"])
            L.append("")
        for t in p.get("points", []):
            sids = "、".join(t.get("sids") or []) or "（人工补充）"
            L.append("- %s（素材：%s）" % (t.get("text", ""), sids))
            if t.get("quote"):
                L.append("  - 原文：%s" % t["quote"])
            if t.get("note"):
                L.append("  - 备注：%s" % t["note"])
        if p.get("speaker_note"):
            L.append("")
            L.append("> 讲稿：%s" % p["speaker_note"])
        L.append("")
    return "\n".join(L)


def render_docx(ol, mats, out):
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        die("未安装 python-docx，请先在隔离环境执行：pip install python-docx")

    mi = material_index(mats)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = CN_FONT
    st.element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", CN_FONT)

    meta = ol.get("meta", {}) or {}

    
    _para(doc, meta.get("name", "对外演示材料"), size=22, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    _para(doc, "演示脚本（含大纲、页面要点与讲稿备注）", size=12,
          color=(0x88, 0x88, 0x88), align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    info = [("受众", meta.get("audience") or "（待填）"),
            ("演示目标", meta.get("goal") or "（待填）"),
            ("计划时长", "%s 分钟" % (meta.get("duration_min") or "（待填）")),
            ("页面数", "%d 页" % len(ol.get("pages", []))),
            ("素材条目", "%d 条" % len(mats)),
            ("生成时间", meta.get("created_at") or datetime.datetime.now().strftime("%Y-%m-%d %H:%M"))]
    if meta.get("edited_by"):
        info.append(("人工调整", "%s（%s）" % (meta["edited_by"], meta.get("edit_note", ""))))
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in info:
        row = t.add_row().cells
        row[0].text = ""
        _set_font(row[0].paragraphs[0].add_run(k), size=10, bold=True)
        row[1].text = ""
        _set_font(row[1].paragraphs[0].add_run(str(v)), size=10)

    
    doc.add_page_break()
    _para(doc, "一、演示大纲总览", size=16, bold=True, space_after=10)
    tb = doc.add_table(rows=1, cols=6)
    tb.style = "Table Grid"
    hdr = ["页码", "角色", "标题", "要点数", "时长(分)", "素材来源"]
    for i, h in enumerate(hdr):
        c = tb.rows[0].cells[i]
        c.text = ""
        _set_font(c.paragraphs[0].add_run(h), size=10, bold=True)
    for p in ol.get("pages", []):
        sids = []
        for pt in p.get("points", []):
            for s in (pt.get("sids") or []):
                if s not in sids:
                    sids.append(s)
        cells = tb.add_row().cells
        vals = [p.get("id", ""), p.get("role", ""), p.get("title", ""),
                str(len(p.get("points", []))), str(p.get("minutes", "")),
                "、".join(sids) or "-"]
        for i, v in enumerate(vals):
            cells[i].text = ""
            _set_font(cells[i].paragraphs[0].add_run(v), size=9)

    
    doc.add_page_break()
    _para(doc, "二、逐页演示脚本", size=16, bold=True, space_after=10)

    for p in ol.get("pages", []):
        _para(doc, "%s  %s" % (p.get("id", ""), p.get("title", "")),
              size=14, bold=True, space_after=2)
        sub = []
        if p.get("role"):
            sub.append("角色：%s" % p["role"])
        if p.get("minutes"):
            sub.append("建议时长：%s 分钟" % p["minutes"])
        if sub:
            _para(doc, "（%s）" % " / ".join(sub), size=9,
                  color=(0x88, 0x88, 0x88), space_after=4)
        if p.get("goal"):
            _para(doc, "目标：%s" % p["goal"], size=10, italic=True, space_after=6)

        pts = p.get("points", [])
        if not pts:
            _para(doc, "（本页暂无要点，可人工补充）", size=10,
                  color=(0x99, 0x99, 0x99), space_after=6)
        for pt in pts:
            sids = pt.get("sids") or []
            tag = "、".join(sids) if sids else "人工补充"
            _para(doc, "• %s" % pt.get("text", ""), size=11, space_after=2)
            if pt.get("quote"):
                _para(doc, "  原文（%s）：%s" % (tag, pt["quote"]), size=9,
                      color=(0x77, 0x77, 0x77), italic=True, space_after=2)
            if pt.get("note"):
                _para(doc, "  说明：%s" % pt["note"], size=9,
                      color=(0x99, 0x66, 0x00), space_after=2)
        if p.get("speaker_note"):
            _para(doc, "讲稿提示：%s" % p["speaker_note"], size=10,
                  color=(0x44, 0x66, 0x99), space_after=8)
        else:
            _para(doc, "", size=6, space_after=8)

    
    doc.add_page_break()
    _para(doc, "三、附录：素材溯源表", size=16, bold=True, space_after=8)
    _para(doc, "每条要点的引文均逐字取自下表素材，编号一一对应，便于人工复核与调整。",
          size=9, color=(0x88, 0x88, 0x88), space_after=8)
    mt = doc.add_table(rows=1, cols=4)
    mt.style = "Table Grid"
    for i, h in enumerate(["编号", "来源 / 章节", "原文（节选）", "引用页"]):
        c = mt.rows[0].cells[i]
        c.text = ""
        _set_font(c.paragraphs[0].add_run(h), size=10, bold=True)

    used = {}
    for p in ol.get("pages", []):
        for pt in p.get("points", []):
            for s in (pt.get("sids") or []):
                used.setdefault(s, [])
                if p.get("id") not in used[s]:
                    used[s].append(p.get("id"))
    for m in mats:
        cells = mt.add_row().cells
        vals = [m.get("id", ""),
                "%s / %s" % (m.get("source", ""), m.get("section", "")),
                (m.get("text", "") or "").replace("\n", " ")[:120],
                "、".join(used.get(m.get("id", ""), [])) or "未引用"]
        for i, v in enumerate(vals):
            cells[i].text = ""
            _set_font(cells[i].paragraphs[0].add_run(v), size=8)

    d = os.path.dirname(os.path.abspath(out))
    if d and not os.path.isdir(d):
        os.makedirs(d)
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description="大纲 → Word 演示脚本")
    ap.add_argument("outline", help="outline.json")
    ap.add_argument("--materials", help="materials.json（默认取大纲内记录的路径）")
    ap.add_argument("-o", "--out", help="输出 docx")
    ap.add_argument("--md", help="同时输出 Markdown 大纲（便于人工调整）")
    ap.add_argument("--reindex", action="store_true", help="按当前顺序重排页面/要点编号")
    ap.add_argument("--out-json", help="--reindex 时另存的大纲文件")
    args = ap.parse_args()

    if not os.path.isfile(args.outline):
        die("大纲文件不存在：%s" % args.outline)
    ol = load_json(args.outline)
    mats = []
    mpath = args.materials or (ol.get("meta", {}) or {}).get("materials", "")
    if mpath and os.path.isfile(mpath):
        mats = load_json(mpath).get("materials", [])
    elif isinstance(ol.get("materials"), list):
        mats = ol["materials"]

    if args.reindex:
        ol = reindex(ol)
        target = args.out_json or args.outline
        save_json(target, ol)
        print("已重排编号并保存：%s" % target)

    if args.md:
        d = os.path.dirname(os.path.abspath(args.md))
        if d and not os.path.isdir(d):
            os.makedirs(d)
        with open(args.md, "w", encoding="utf-8") as f:
            f.write(render_md(ol, mats))
        print("Markdown 大纲已写入：%s" % args.md)

    if args.out:
        render_docx(ol, mats, args.out)
        print("Word 演示脚本已生成：%s" % args.out)
        print("  页面 %d 页 | 素材 %d 条 | 要点 %d 条"
              % (len(ol.get("pages", [])), len(mats),
                 sum(len(p.get("points", [])) for p in ol.get("pages", []))))
    return 0


if __name__ == "__main__":
    sys.exit(main())
